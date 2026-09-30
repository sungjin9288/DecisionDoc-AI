"""Versioned applicability contract and lossless portable renderers."""
from __future__ import annotations

import hashlib
import html
import io
import json
import zipfile
from collections.abc import Mapping, Sequence
from typing import Any

from app.schemas.procurement_applicability_export import ProcurementApplicabilityExport
from app.schemas.procurement_binding import ProcurementSourceBinding
from app.services.procurement_decision_package.package_constants import INCLUDED_ARTIFACT_ORDER


PACKAGE_SCHEMA_PURPOSE_V2 = "procurement_decision_package.v2"
APPLICABILITY_DOCX_NAME = "requirement_applicability.docx"
V3_ARTIFACT_ORDER = (*INCLUDED_ARTIFACT_ORDER, APPLICABILITY_DOCX_NAME)


def artifact_order(package_doc: Mapping[str, Any]) -> Sequence[str]:
    return V3_ARTIFACT_ORDER if package_doc.get("schema_purpose") == PACKAGE_SCHEMA_PURPOSE_V2 else INCLUDED_ARTIFACT_ORDER


def validate_applicability(value: dict, binding: ProcurementSourceBinding | None = None) -> dict:
    projection = ProcurementApplicabilityExport.model_validate(value)
    if binding is not None:
        fingerprints = [item.model_dump(mode="json") for item in binding.snapshots]
        source_hash = hashlib.sha256(json.dumps(
            fingerprints, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False,
        ).encode("utf-8")).hexdigest()
        if (projection.decision_id, projection.decision_revision, projection.source_set_sha256) != (
            binding.decision_id, binding.decision_revision, source_hash,
        ):
            raise ValueError("requirement applicability source_binding mismatch")
        sources = {item.snapshot_id: item.sha256 for item in binding.snapshots}
        for row in projection.requirements:
            if not row.stale and sources.get(row.snapshot_id) != row.snapshot_sha256:
                raise ValueError("current requirement snapshot is not in source_binding")
    return projection.model_dump(mode="json")


def applicability_counts(projection: dict) -> dict[str, int]:
    rows = projection["requirements"]
    return {
        "total": len(rows),
        **{status: sum(row["applicability"] == status for row in rows)
           for status in ("applies", "not_applicable", "unknown")},
        "stale": sum(row["stale"] for row in rows),
    }


def applicability_blocks(projection: dict) -> list[tuple[str, str]]:
    """A shared semantic render plan; literal values never enter Markdown parsing."""
    projection = validate_applicability(projection)
    counts = applicability_counts(projection)
    blocks = [("h1", "요구사항 적용 여부 검토"),
              ("p", "사람이 기록한 판단입니다. 상위 범주 상태와 점수는 변경하지 않습니다."),
              ("p", "원문 bytes 재검증 안 함 (source_bytes_verified: false)"),
              ("p", "운영 승인 아님 (operational_approval: false)"),
              ("h2", "세부 요구사항 집계"),
              ("p", f"전체 {counts['total']} · 적용 {counts['applies']} · 해당 없음(N/A) {counts['not_applicable']} · 미확인 {counts['unknown']} · 근거 변경(stale) {counts['stale']}"),
              ("p", "stale은 미확인에 포함됩니다. 부모 범주의 준비 상태와 별도 집계입니다."),
              ("source", f"형식: {projection['schema_version']}"),
              ("source", f"공고 판단 ID: {projection['decision_id']}"),
              ("source", f"판단 revision: {projection['decision_revision']}"),
              ("source", f"현재 source-set SHA-256: {projection['source_set_sha256']}")]
    labels = {"applies": "적용", "not_applicable": "해당 없음(N/A)", "unknown": "미확인"}

    def source(item: dict[str, Any]) -> None:
        blocks.extend([
            ("source", f"Snapshot ID: {item['snapshot_id']}"),
            ("source", f"Snapshot SHA-256: {item['snapshot_sha256']}"),
            ("source", f"Source-set SHA-256: {item['source_set_sha256']}"),
            ("source", f"인용 구간 (Unicode code point): [{item['quote_start']}, {item['quote_end']})"),
        ])

    for index, row in enumerate(projection["requirements"], 1):
        blocks.extend([
            ("h2", f"요구사항 {index}. {row['title']}"),
            ("source", f"요구사항 ID: {row['requirement_id']}"),
            ("p", f"범주: {row['category']}"),
            ("p", f"현재 적용 여부: {labels[row['applicability']]} ({row['applicability']})"),
            ("p", "근거 상태: 변경됨 (stale: true)" if row["stale"] else "근거 상태: 동일 source set (stale: false)"),
            ("p", f"작성 시각: {row['created_at']}"),
            ("h3", "원문 인용"), ("quote", row["quote"]),
        ])
        source(row)
        blocks.append(("p", f"판단 이력: {len(row['annotations'])}건"))
        for event, annotation in enumerate(row["annotations"], 1):
            blocks.extend([
                ("h3", f"이력 {event}. {labels[annotation['applicability']]} ({annotation['applicability']})"),
                ("p", f"기록 시각: {annotation['created_at']}"),
                ("p", "판단 사유"), ("quote", annotation["rationale"]),
                ("p", "당시 원문 인용"), ("quote", annotation["quote"]),
            ])
            source(annotation)
            blocks.extend([
                ("source", f"주석 ID: {annotation['annotation_id']}"),
                ("source", f"Expected decision revision: {annotation['expected_decision_revision']}"),
                ("source", f"Operation ID: {annotation['operation_id']}"),
            ])
    return blocks


def applicability_lines(projection: dict) -> list[str]:
    return [text for _, text in applicability_blocks(projection)]


def render_applicability_markdown(projection: dict) -> str:
    sections = []
    for kind, text in applicability_blocks(projection):
        # Headings include untrusted titles. Numeric entities prevent Markdown/HTML
        # interpretation without altering the visible Unicode text.
        if kind.startswith("h"):
            safe = "".join(f"&#{ord(char)};" if char in "\\`*_{}[]()<>#+!|&\r\n" else char
                           for char in text)
            sections.append("#" * (int(kind[1]) + 1) + " " + safe)
        else:
            sections.append("\n".join("    " + line for line in text.split("\n")))
    return "\n\n" + "\n\n".join(sections)


def render_applicability_html(projection: dict) -> str:
    parts = ['<section data-requirement-applicability>']
    for kind, text in applicability_blocks(projection):
        tag = {"h1": "h2", "h2": "h3", "h3": "h4", "quote": "blockquote"}.get(kind, "p")
        parts.append(f'<{tag} class="applicability-{kind}" style="white-space:pre-wrap;overflow-wrap:anywhere">'
                     + html.escape(text, quote=True) + f'</{tag}>')
    return "".join(parts) + "</section>"


def build_applicability_docx(projection: dict) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    from app.services.docx_service import build_docx, normalize_zip_metadata

    # The shared Markdown parser transforms inline syntax. Use its document shell,
    # then native runs so untrusted quotes remain exact, editable text.
    document = Document(io.BytesIO(build_docx([], title="Requirement Applicability")))
    # Reuse page/font configuration, but not the generic submission cover copy.
    document._body.clear_content()
    title_properties = document.styles["Title"].element.pPr
    if title_properties is not None:
        for border in list(title_properties.findall(qn("w:pBdr"))):
            title_properties.remove(border)
    for kind, text in applicability_blocks(projection):
        # Word represents all line endings as breaks; the JSON keeps exact source bytes.
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        if kind == "h1":
            paragraph = document.add_paragraph(text, style="Title")
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.size = Pt(20)
        elif kind.startswith("h"):
            document.add_heading(text, level=int(kind[1]))
        else:
            paragraph = document.add_paragraph(text, style="Quote" if kind == "quote" else None)
            paragraph.paragraph_format.space_after = Pt(4 if kind == "source" else 7)
            for run in paragraph.runs:
                run.font.italic = False
                if kind == "source":
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor.from_string("4B5563")
    output = io.BytesIO()
    document.save(output)
    return normalize_zip_metadata(output.getvalue())


def validate_applicability_docx(content: bytes, projection: dict) -> None:
    expected = build_applicability_docx(projection)
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as actual, zipfile.ZipFile(io.BytesIO(expected)) as reference:
            if actual.namelist() != reference.namelist():
                raise ValueError("requirement applicability DOCX members mismatch")
            # Compare regenerated OOXML, not attacker-rehashed archive fingerprints.
            for info in actual.infolist():
                target = reference.read(info.filename)
                if info.file_size != len(target) or actual.read(info) != target:
                    raise ValueError("requirement applicability DOCX content mismatch")
    except zipfile.BadZipFile as exc:
        raise ValueError("invalid requirement applicability DOCX") from exc
