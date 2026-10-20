"""Read RFPs and drafts into located text blocks, without truncation.

The attachment extractors used by the web app cap text at 12,000 characters and
drop structure. Requirement tracing needs the whole document and where each line
sits (heading path, PDF page), so these readers keep both.
"""
from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

MAX_INPUT_BYTES = 100 * 1024 * 1024
MAX_ZIP_MEMBERS = 2000
MAX_ZIP_UNCOMPRESSED_BYTES = 300 * 1024 * 1024
SUPPORTED_SUFFIXES = (".md", ".txt", ".docx", ".hwpx", ".pdf")
_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(\S.*)$")
_FENCE_RE = re.compile(r"^\s*(```|~~~)")
_STYLE_LEVEL_RE = re.compile(r"(?:heading|제목|개요)\s*(\d)", re.IGNORECASE)
# DecisionDoc's HWPX writer styles table header rows as 제목3 with this prefix.
_HWPX_TABLE_LABEL = "표:"


class ReaderError(ValueError):
    """The file is missing, too large, unsupported or malformed."""


@dataclass(frozen=True)
class Block:
    text: str
    heading_level: int | None = None
    page: int | None = None


@dataclass(frozen=True)
class SourceDocument:
    path: str
    format: str
    blocks: tuple[Block, ...]

    @property
    def has_pages(self) -> bool:
        return any(block.page is not None for block in self.blocks)

    def full_text(self) -> str:
        return "\n".join(block.text for block in self.blocks)

    def heading_paths(self) -> list[str]:
        """The heading path (``상위 > 하위``) in force at each block."""
        stack: list[tuple[int, str]] = []
        paths: list[str] = []
        for block in self.blocks:
            if block.heading_level is not None:
                stack = [item for item in stack if item[0] < block.heading_level]
                stack.append((block.heading_level, block.text))
            paths.append(" > ".join(text for _, text in stack))
        return paths


def read_document(path: str | Path) -> SourceDocument:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ReaderError(f"지원하지 않는 형식입니다: {source.name} (지원: {', '.join(SUPPORTED_SUFFIXES)})")
    if not source.is_file():
        raise ReaderError(f"파일이 없습니다: {source}")
    if source.stat().st_size > MAX_INPUT_BYTES:
        raise ReaderError(f"파일이 너무 큽니다: {source.name}")
    raw = source.read_bytes()
    readers = {".md": _read_markdown, ".txt": _read_plain, ".docx": _read_docx, ".hwpx": _read_hwpx, ".pdf": _read_pdf}
    blocks = readers[suffix](raw)
    return SourceDocument(path=str(source), format=suffix.lstrip("."), blocks=tuple(blocks))


def _decode(raw: bytes) -> str:
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ReaderError("UTF-8 텍스트가 아닙니다") from exc


def _read_plain(raw: bytes) -> list[Block]:
    return [Block(line.strip()) for line in _decode(raw).splitlines() if line.strip()]


def _read_markdown(raw: bytes) -> list[Block]:
    blocks: list[Block] = []
    in_fence = False
    for line in _decode(raw).splitlines():
        if _FENCE_RE.match(line):
            in_fence = not in_fence
            continue
        if not line.strip():
            continue
        match = None if in_fence else _MD_HEADING_RE.match(line.strip())
        if match:
            blocks.append(Block(match.group(2).strip(), heading_level=len(match.group(1))))
        else:
            blocks.append(Block(line.strip()))
    return blocks


def _style_level(style_name: str) -> int | None:
    match = _STYLE_LEVEL_RE.search(style_name or "")
    return int(match.group(1)) if match else None


def _read_docx(raw: bytes) -> list[Block]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(BytesIO(raw))
    except Exception as exc:  # python-docx raises several types for broken packages
        raise ReaderError(f"DOCX를 열 수 없습니다: {exc}") from exc
    blocks: list[Block] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if text:
                blocks.append(Block(text, heading_level=_style_level(paragraph.style.name if paragraph.style else "")))
        elif tag == "tbl":
            for row in Table(child, document).rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append(Block(" | ".join(cells)))
    return blocks


def _open_zip(raw: bytes) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(BytesIO(raw))
    except zipfile.BadZipFile as exc:
        raise ReaderError("ZIP 형식이 아닙니다") from exc
    infos = archive.infolist()
    if len(infos) > MAX_ZIP_MEMBERS or sum(info.file_size for info in infos) > MAX_ZIP_UNCOMPRESSED_BYTES:
        raise ReaderError("압축 해제 크기 또는 항목 수가 한도를 넘습니다")
    return archive


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _hwpx_style_names(archive: zipfile.ZipFile) -> dict[str, str]:
    from defusedxml import ElementTree

    if "Contents/header.xml" not in archive.namelist():
        return {}
    root = ElementTree.fromstring(archive.read("Contents/header.xml"))
    return {el.get("id", ""): el.get("name", "") for el in root.iter() if _local(el.tag) == "style"}


def _hwpx_paragraph_text(paragraph) -> str:  # noqa: ANN001 - ElementTree element
    # Only this paragraph's own runs; table cells hold nested paragraphs read separately.
    parts = [
        node.text or ""
        for run in paragraph
        if _local(run.tag) == "run"
        for node in run
        if _local(node.tag) == "t"
    ]
    return "".join(parts).strip()


def _read_hwpx(raw: bytes) -> list[Block]:
    from defusedxml import ElementTree

    archive = _open_zip(raw)
    styles = _hwpx_style_names(archive)
    sections = sorted(
        (name for name in archive.namelist() if re.fullmatch(r"Contents/section\d+\.xml", name)),
        key=lambda name: int(re.search(r"\d+", name).group()),
    )
    if not sections:
        raise ReaderError("HWPX 본문(Contents/section*.xml)이 없습니다")
    blocks: list[Block] = []
    for name in sections:
        root = ElementTree.fromstring(archive.read(name))
        for paragraph in root.iter():
            if _local(paragraph.tag) != "p":
                continue
            text = _hwpx_paragraph_text(paragraph)
            if not text:
                continue
            level = _style_level(styles.get(paragraph.get("styleIDRef", ""), ""))
            if text.startswith(_HWPX_TABLE_LABEL):
                level = None
            blocks.append(Block(text, heading_level=level))
    return blocks


def _pdf_page_texts(raw: bytes) -> list[str]:
    # pdfplumber is the declared dependency; PyMuPDF is accepted where it is installed instead.
    try:
        import pdfplumber
    except ImportError:
        try:
            import fitz
        except ImportError as exc:
            raise ReaderError("PDF를 읽으려면 pdfplumber가 필요합니다") from exc
        fitz.TOOLS.mupdf_display_errors(False)
        with fitz.open(stream=raw, filetype="pdf") as document:
            return [page.get_text() for page in document]
    with pdfplumber.open(BytesIO(raw)) as document:
        return [page.extract_text() or "" for page in document.pages]


def _read_pdf(raw: bytes) -> list[Block]:
    try:
        pages = _pdf_page_texts(raw)
    except ReaderError:
        raise
    except Exception as exc:  # parser-specific errors for broken files
        raise ReaderError(f"PDF를 열 수 없습니다: {exc}") from exc
    return [
        Block(line.strip(), page=number)
        for number, text in enumerate(pages, start=1)
        for line in text.splitlines()
        if line.strip()
    ]


def pdf_page_count(raw: bytes) -> int:
    try:
        return len(_pdf_page_texts(raw))
    except ReaderError:
        raise
    except Exception as exc:  # parser-specific errors for broken files
        raise ReaderError(f"PDF를 열 수 없습니다: {exc}") from exc


def metadata_text(path: str | Path) -> str:
    """Author, company and other properties a blind evaluation must not reveal."""
    source = Path(path)
    raw = source.read_bytes()
    suffix = source.suffix.lower()
    if suffix == ".pdf":
        try:
            import pdfplumber
        except ImportError:
            try:
                import fitz
            except ImportError as exc:
                raise ReaderError("PDF를 읽으려면 pdfplumber가 필요합니다") from exc
            with fitz.open(stream=raw, filetype="pdf") as document:
                return " ".join(str(value) for value in (document.metadata or {}).values() if value)
        with pdfplumber.open(BytesIO(raw)) as document:
            return " ".join(str(value) for value in (document.metadata or {}).values() if value)
    if suffix in (".docx", ".hwpx"):
        archive = _open_zip(raw)
        parts = [
            name for name in archive.namelist()
            if name.startswith(("docProps/", "META-INF/")) or name.endswith(("content.hpf", "settings.xml"))
        ]
        return " ".join(archive.read(name).decode("utf-8", "ignore") for name in parts)
    return ""


def read_documents(paths: list[str | Path]) -> SourceDocument:
    """Read several drafts as one document in the given order (e.g. a proposal bundle's four parts)."""
    documents = [read_document(path) for path in paths]
    if len(documents) == 1:
        return documents[0]
    return SourceDocument(
        path=", ".join(document.path for document in documents),
        format="+".join(document.format for document in documents),
        blocks=tuple(block for document in documents for block in document.blocks),
    )
