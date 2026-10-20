"""package_cover=False drops the DecisionDoc cover and per-document cards, keeping the body."""
from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document

from app.schemas.visual_assets import GovDocOptions
from app.services.docx_service import build_docx
from app.services.hwp_service import build_hwp
from tests.async_helper import run_async

DOCS = [
    {"doc_type": "proposal_body", "markdown": "# Ⅰ. 사업 이해\n\n## 1. 추진 배경\n결재 완료 문서를 자동 보관한다.\n"},
    {"doc_type": "proposal_tech", "markdown": "# Ⅱ. 기술 부문\n\n## 1. 통합 검색\nSFR-002 통합 검색을 제공한다.\n"},
]
PACKAGE_MARKERS = ("DecisionDoc AI Export", "완성형 문서 패키지", "문서 01 / 02", "구성 지표")


def _docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([paragraph.text for paragraph in document.paragraphs] + cells)


def _hwpx_text(content: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return archive.read("Contents/section0.xml").decode("utf-8")


def _pdf_text(content: bytes) -> str:
    from app.public_sector.readers import _pdf_page_texts

    return "\n".join(_pdf_page_texts(content))


def _render(fmt: str, options: GovDocOptions | None) -> str:
    if fmt == "docx":
        return _docx_text(build_docx(DOCS, title="[합성] 정성 제안서", gov_options=options))
    if fmt == "hwpx":
        return _hwpx_text(build_hwp(DOCS, title="[합성] 정성 제안서", gov_options=options))
    from app.services.pdf_service import build_pdf

    return _pdf_text(run_async(build_pdf(DOCS, title="[합성] 정성 제안서", gov_options=options)))


@pytest.mark.parametrize("fmt", ["docx", "hwpx", "pdf"])
def test_package_cover_is_kept_by_default(fmt: str) -> None:
    text = _render(fmt, None)

    assert "완성형 문서 패키지" in text


@pytest.mark.parametrize("fmt", ["docx", "hwpx", "pdf"])
def test_package_cover_false_keeps_only_the_body(fmt: str) -> None:
    text = _render(fmt, GovDocOptions(package_cover=False))

    assert not [marker for marker in PACKAGE_MARKERS if marker in text]
    assert "결재 완료 문서를 자동 보관한다." in text and "SFR-002 통합 검색을 제공한다." in text
