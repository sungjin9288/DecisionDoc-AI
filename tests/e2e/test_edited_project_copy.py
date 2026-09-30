import json
from uuid import uuid4
import zipfile

import pytest
from docx import Document
from playwright.sync_api import expect

from tests.e2e.test_main_flow import _generate_to_results


@pytest.mark.parametrize(
    "viewport", [{"width": 1280, "height": 900}, {"width": 390, "height": 844}]
)
def test_edited_project_copy_save_reload_download(
    page, live_server, tmp_path, viewport
):
    page_errors = []
    page.on('pageerror', lambda error: page_errors.append(str(error)))
    page.set_viewport_size(viewport)
    page.route(
        "**/*",
        lambda route: (
            route.continue_()
            if route.request.url.startswith(live_server["base_url"] + "/")
            else route.abort()
        ),
    )
    project = page.evaluate(
        """async name => {
        const response = await fetch('/projects', {method:'POST', headers:{...getAuthHeaders(), 'Content-Type':'application/json'},
          body:JSON.stringify({name})});
        if (!response.ok) throw new Error(await response.text());
        return response.json();
    }""",
        "Copy lifecycle " + str(uuid4()),
    )
    page.reload()
    page.locator(".bundle-card").first.click()
    page.select_option("#project-select", project["project_id"])
    _generate_to_results(
        page, "Edited copy lifecycle", "Verify local saved document persistence"
    )
    # Project auto-link completes after the generation completion event.
    page.wait_for_function(
        """async projectId => {
        const response = await fetch('/projects/' + projectId, {headers:getAuthHeaders()});
        return (await response.json()).documents.length === 1;
    }""",
        arg=project["project_id"],
    )
    original = page.evaluate(
        """async id => (await (await fetch('/projects/'+id,
        {headers:getAuthHeaders()})).json()).documents[0]""",
        project["project_id"],
    )
    page.click("#edit-btn")
    edited_markdown = (
        "## Draft section\n\nSaved edited copy marker\n\n"
        "| Kind | Path | Options |\n| --- | --- | --- |\n"
        r"| File | C:\docs\d.docx | A \| B |"
    )
    page.locator("#doc-pane .doc-pane-content").fill(edited_markdown)
    generation_requests = []
    page.on('request', lambda request: generation_requests.append(request.url)
            if request.method == 'POST' and '/generate/' in request.url else None)
    with page.expect_download() as markdown_download:
        page.click('#export-btn')
    markdown_path = tmp_path / 'live-draft.md'
    markdown_download.value.save_as(markdown_path)
    assert markdown_path.read_bytes() == edited_markdown.encode('utf-8')
    assert generation_requests == []
    page.locator("#doc-pane .doc-pane-content").scroll_into_view_if_needed()
    page.screenshot(path=str(tmp_path / "edited-copy-source.png"))
    with page.expect_download() as live_download:
        page.locator('[data-result-export="docx"]').click()
    live_path = tmp_path / "live-draft.docx"
    live_download.value.save_as(live_path)
    live_export = Document(live_path)
    assert any(p.text == 'Saved edited copy marker' for p in live_export.paragraphs)
    live_table = next(table for table in live_export.tables if table.cell(0, 0).text == 'Kind')
    assert [cell.text for cell in live_table.rows[1].cells] == ['File', r'C:\docs\d.docx', 'A | B']
    expect(page.locator('#doc-pane textarea')).to_have_value(edited_markdown)
    page.click('#edit-btn')
    rewritten = edited_markdown.split('\n\n', 1)[1].replace(
        'Saved edited copy marker', 'Saved edited copy marker after section rewrite'
    )
    page.route('**/generate/rewrite-section', lambda route: route.fulfill(
        json={'rewritten': rewritten}, status=200
    ))
    page.locator('.section-rewrite-btn').click()
    page.locator('#rewrite-instruction').fill('Keep the table and clarify the text')
    page.screenshot(path=str(tmp_path / 'section-rewrite-dialog.png'))
    with page.expect_request('**/generate/rewrite-section') as rewrite_request:
        page.click('#rewrite-confirm')
    assert rewrite_request.value.post_data_json['current_content'] == edited_markdown.split('\n\n', 1)[1]
    edited_markdown = '## Draft section\n' + rewritten + '\n'
    expect(page.locator('.rewrite-modal-overlay')).to_have_count(0)
    expect(page.locator('#doc-pane')).to_contain_text('after section rewrite')
    with page.expect_response(
        lambda response: response.url.endswith("/edited-copies")
    ) as saved:
        page.click("#save-edited-copy-btn")
    assert saved.value.status == 200, saved.value.text()
    copy = saved.value.json()["document"]
    page.reload()
    page.evaluate("switchPage('project-page')")
    page.locator(f'[data-project-open="{project["project_id"]}"]').click()
    page.locator(
        f'[data-project-detail-action="doc-open-edited"][data-doc-id="{copy["doc_id"]}"]'
    ).click()
    expect(page.locator("#doc-pane")).to_contain_text("Saved edited copy marker")
    expect(page.locator("#doc-pane")).to_contain_text("after section rewrite")
    generation_requests.clear()
    with page.expect_download() as reopened_download:
        page.click('#export-btn')
    reopened_path = tmp_path / 'reopened-draft.md'
    reopened_download.value.save_as(reopened_path)
    assert reopened_path.read_bytes() == edited_markdown.encode('utf-8')
    assert generation_requests == []
    expect(page.locator('#doc-pane .section-rewrite-btn')).to_have_count(1)
    rewrite_button = page.locator('#doc-pane .section-rewrite-btn')
    if viewport['width'] < 768:
        expect(rewrite_button).to_have_css('opacity', '1')
    rewrite_button.focus()
    expect(rewrite_button).to_have_css('opacity', '1')
    page.locator('#doc-pane .section-rewrite-btn').click()
    expect(page.locator('.rewrite-modal')).to_be_visible()
    page.click('#rewrite-cancel')
    expect(page.locator("#doc-pane")).to_contain_text(r"C:\docs\d.docx")
    expect(page.locator("#doc-pane")).to_contain_text("A | B")
    expect(page.locator('#cost-meta')).to_have_text('저장된 편집본 · AI 재생성 없음')
    expect(page.locator("[data-result-export-zip]")).to_have_count(0)
    page.evaluate("""async () => {
        await Promise.all(document.getElementById('results').getAnimations().map(animation => animation.finished));
    }""")
    page.locator('#doc-pane').scroll_into_view_if_needed()
    page.screenshot(path=str(tmp_path / "edited-copy-reopened.png"))
    overflow = page.evaluate("""() => [...document.querySelectorAll('main *')]
        .filter(el => el.getBoundingClientRect().right > innerWidth + 1)
        .slice(0, 12).map(el => ({tag:el.tagName,id:el.id,class:el.className,
            width:el.getBoundingClientRect().width,right:el.getBoundingClientRect().right}))""")
    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth"), overflow
    page.evaluate("switchPage('project-page')")
    with page.expect_download() as download:
        page.locator(
            f'[data-project-detail-action="doc-download"][data-doc-id="{copy["doc_id"]}"][data-format="docx"]'
        ).click()
    path = tmp_path / "copy.docx"
    download.value.save_as(path)
    with zipfile.ZipFile(path) as archive:
        assert b"Saved edited copy marker" in archive.read("word/document.xml")
    exported = Document(path)
    table = next(table for table in exported.tables if table.cell(0, 0).text == "Kind")
    assert [cell.text for cell in table.rows[1].cells] == ["File", r"C:\docs\d.docx", "A | B"]
    restored = page.evaluate(
        """async id => (await (await fetch('/projects/'+id,
        {headers:getAuthHeaders()})).json()).documents""",
        project["project_id"],
    )
    assert len(restored) == 2
    assert restored[0]["doc_snapshot"] == original["doc_snapshot"]
    assert (
        json.loads(restored[1]["doc_snapshot"])[0]["markdown"]
        == edited_markdown
    )
    page.evaluate("switchPage('generate')")
    page.click('#edit-btn')
    page.locator('#doc-pane .doc-pane-content').fill('Second saved revision')
    page.click('#edit-btn')
    with page.expect_response(lambda response: response.url.endswith('/edited-copies')) as successor:
        page.click('#save-edited-copy-btn')
    assert successor.value.status == 200
    assert successor.value.json()['document']['edited_copy']['parent_id'] == copy['doc_id']
    assert page_errors == []
