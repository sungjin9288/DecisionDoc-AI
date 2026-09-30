"""Geometry checks for text on PPTX cover, agenda and section-divider slides.

PowerPoint does not wrap a python-pptx text box unless ``word_wrap`` is set, so
long summaries used to run past their cards and the slide edge. Divider and
cover cards also overlapped the title/subtitle placeholders.
"""
from io import BytesIO
import math

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.services.pptx_service import build_pptx_from_docs

_LONG_LEAD = (
    "본 사업은 여러 기관에 흩어진 행정 데이터를 하나의 통합 플랫폼으로 모아 "
    "정책 담당자가 실시간 지표를 확인하도록 지원하는 것을 목표로 합니다."
)
_DOCS = [
    {"doc_type": "adr", "markdown": f"# 공공 데이터 통합 플랫폼 구축 결정\n\n{_LONG_LEAD}\n\n## 배경\n\n현재 데이터는 분산되어 있습니다."},
    {"doc_type": "ops_checklist", "markdown": "# 운영 점검표\n\n운영 조직은 장애 대응 절차와 백업 복구 훈련 결과를 분기마다 점검하고 개선 과제를 관리 대장에 기록해야 합니다.\n"},
]
_EMU_PER_INCH = 914400


def _deck() -> Presentation:
    return Presentation(BytesIO(build_pptx_from_docs(_DOCS, "긴 요약 배치 확인")))


def _box(shape) -> tuple[float, float, float, float]:
    return (
        shape.left / _EMU_PER_INCH,
        shape.top / _EMU_PER_INCH,
        (shape.left + shape.width) / _EMU_PER_INCH,
        (shape.top + shape.height) / _EMU_PER_INCH,
    )


def _overlaps(a, b) -> bool:
    return a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]


def _cards(slide):
    return [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]


def _em_width(text: str) -> float:
    return sum(1.0 if ord(char) >= 0x1100 else 0.6 for char in text)


def test_cover_and_divider_text_placeholders_do_not_sit_under_cards():
    deck = _deck()
    divider = next(
        slide for slide in deck.slides
        if slide.shapes.title is not None and slide.shapes.title.text == "공공 데이터 통합 플랫폼 구축 결정"
    )
    for slide in (deck.slides[0], divider):
        placeholders = [shape for shape in slide.placeholders if shape.text_frame.text.strip()]
        assert len(placeholders) == 2
        for placeholder in placeholders:
            left, _, right, _ = _box(placeholder)
            assert left >= 0.5 and right - left >= 6.0, (placeholder.text_frame.text, left, right)
            for card in _cards(slide):
                assert not _overlaps(_box(placeholder), _box(card)), placeholder.text_frame.text


def test_card_text_wraps_and_fits_inside_its_card():
    deck = _deck()
    checked = 0
    for slide in list(deck.slides)[:4]:
        cards = [_box(card) for card in _cards(slide)]
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.TEXT_BOX:
                continue
            box = _box(shape)
            card = next(card for card in cards if card[0] <= box[0] and box[2] <= card[2] and card[1] <= box[1] < card[3])
            frame = shape.text_frame
            assert frame.word_wrap is True, frame.text
            font_pt = frame.paragraphs[0].runs[0].font.size.pt
            per_line = (box[2] - box[0] - 0.2) * 72 / font_pt
            lines = sum(max(1, math.ceil(_em_width(p.text) / per_line)) for p in frame.paragraphs)
            text_bottom = box[1] + 0.05 + lines * font_pt * 1.2 / 72
            assert text_bottom <= card[3] + 0.01, (frame.text, text_bottom, card)
            checked += 1
    assert checked >= 8


def test_agenda_card_keeps_document_label_and_marks_shortened_lead():
    deck = _deck()
    agenda_texts = [shape.text_frame.text for shape in deck.slides[1].shapes if shape.has_text_frame]

    body = next(text for text in agenda_texts if text.startswith("기술 의사결정 기록 (ADR)\n"))
    lead = body.split("\n", 1)[1]
    assert lead.endswith("…")
    assert _LONG_LEAD.startswith(lead[:-1].rstrip())


_STRUCTURED_ITEMS = [
    {
        "page": 1, "title": "교차로 안전 AI 적용 방안",
        "core_message": "우회전 일시정지 감지와 보행자 보호를 우선 적용한다.",
        "key_content": "AI 영상 분석으로 위험 이벤트를 조기 탐지하고 운영자 알림을 자동화한다.",
        "evidence_points": ["사고 위험 구간", "운영자 대응 시간", "CCTV 연계 가능성"],
        "visual_type": "의사결정 매트릭스", "visual_brief": "기준별 도입 우선순위 평가표",
        "layout_hint": "좌측 메시지, 우측 매트릭스, 하단 승인 기준",
        "decision_question": "1차 구축 범위를 교차로 안전으로 승인할 것인가?",
        "narrative_role": "대표 승인 전에 사업 범위를 확정하는 판단 장표",
        "content_blocks": ["문제 정의", "도입 범위", "승인 요청"],
        "data_needs": ["사고 위험 구간별 CCTV 현황", "우회전 일시정지 위반 데이터"],
        "acceptance_criteria": ["도입 범위가 명확함", "근거 자료가 장표에 연결됨", "PM 승인 질문에 답함"],
    },
    {
        "page": 2, "title": "통합 관제 플랫폼 단계별 확산 계획",
        "core_message": "1단계 시범 구축 성과를 정량 지표로 검증한 뒤 2단계 전 지역 확산 여부와 예산 배분 비율을 결정한다.",
        "key_content": "시범 구역 12개 교차로의 위험 이벤트 감소율과 운영자 대응 시간 단축을 분기별로 비교한다.",
        "evidence_points": ["시범 구역 위험 이벤트 감소율 분기 비교", "운영자 대응 시간 단축 효과와 인력 재배치 가능성", "시민 민원 변화 추이"],
        "visual_type": "단계별 로드맵", "visual_brief": "시범 구축에서 전 지역 확산까지 분기별 마일스톤과 예산 배분 비율을 함께 표시",
        "layout_hint": "상단 로드맵 타임라인, 하단 좌측 성과 지표, 하단 우측 예산 배분 도넛 차트",
        "decision_question": "시범 성과 기준을 충족하면 2단계 확산 예산을 다음 회계연도에 편성할 것인가?",
        "narrative_role": "확산 여부를 결정하기 위한 조건과 근거를 한 장에 정리하는 승인 장표",
        "content_blocks": ["시범 성과 지표", "확산 조건", "예산 배분 기준", "위험과 대응"],
        "data_needs": ["교차로별 분기 위험 이벤트 로그", "관제 인력 근무 기록"],
        "acceptance_criteria": ["확산 조건이 정량 지표로 정의됨", "예산 배분 근거가 표로 제시됨", "위험 요소별 대응 책임자가 지정됨"],
    },
]


def _structured_deck() -> Presentation:
    from app.services.pptx_service import build_pptx

    deck = build_pptx(
        {"presentation_goal": "단계형 보고서 승인", "slide_outline": _STRUCTURED_ITEMS},
        title="구조화 슬라이드 확인",
        include_outline_overview=True,
    )
    return Presentation(BytesIO(deck))


def test_structured_slides_have_no_negative_shape_sizes():
    # PowerPoint asks to repair a file whose shape extents are negative.
    for slide in _structured_deck().slides:
        for shape in slide.shapes:
            assert shape.width > 0 and shape.height > 0, (shape.name, shape.width, shape.height)


def test_structured_slide_left_cards_hold_their_text_without_overlap():
    deck = _structured_deck()
    structured = [
        slide for slide in deck.slides
        if slide.shapes.title is not None and slide.shapes.title.text in {item["title"] for item in _STRUCTURED_ITEMS}
    ]
    assert len(structured) == 2
    for slide in structured:
        cards = sorted((_box(card) for card in _cards(slide) if _box(card)[0] < 1.0), key=lambda box: box[1])
        assert len(cards) == 4
        for upper, lower in zip(cards, cards[1:]):
            assert upper[3] <= lower[1], (upper, lower)
        assert cards[-1][3] <= 7.3
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.TEXT_BOX:
                continue
            box = _box(shape)
            card = next((card for card in map(_box, _cards(slide)) if card[0] <= box[0] and box[2] <= card[2] and card[1] <= box[1] < card[3]), None)
            if card is None:
                continue
            frame = shape.text_frame
            assert frame.word_wrap is True, frame.text
            font_pt = frame.paragraphs[0].runs[0].font.size.pt
            per_line = (box[2] - box[0] - 0.2) * 72 / font_pt
            lines = sum(max(1, math.ceil(_em_width(p.text) / per_line)) for p in frame.paragraphs)
            assert box[1] + 0.05 + lines * font_pt * 1.2 / 72 <= card[3] + 0.01, (frame.text, card)


def test_structured_slide_keeps_guidance_and_all_content_blocks():
    deck = _structured_deck()
    text = "\n".join(shape.text_frame.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    # content_blocks is limited to three entries by design.
    for block in ("시범 성과 지표", "확산 조건", "예산 배분 기준"):
        assert block in text
    assert "배치 가이드:" in text and "검증 필요: 교차로별 분기 위험 이벤트 로그" in text


def test_structured_slide_without_optional_lists_never_prints_none():
    from app.services.pptx_service import build_pptx

    item = {"page": 1, "title": "현황 분석", "core_message": "현재 업무 흐름을 정리한다.",
            "evidence_points": ["현재 업무 흐름도", "주요 Pain Point"], "visual_type": "프로세스 흐름도"}
    deck = Presentation(BytesIO(build_pptx({"slide_outline": [item]}, title="선택 항목 없음")))
    texts = [shape.text_frame.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame]

    assert "None" not in texts
    assert "근거 블록: 현재 업무 흐름도" in "\n".join(texts)


def test_structured_flow_panel_steps_hold_title_and_text():
    from app.services.pptx_service import build_pptx

    item = {"page": 1, "title": "처리 흐름", "core_message": "세 단계로 처리한다.",
            "evidence_points": ["데이터 수집과 정제", "AI 분석과 판정", "결과 제공과 알림"],
            "visual_type": "프로세스 흐름도"}
    deck = Presentation(BytesIO(build_pptx({"slide_outline": [item]}, title="흐름")))
    slide = next(slide for slide in deck.slides if slide.shapes.title is not None and slide.shapes.title.text == "처리 흐름")
    steps = [shape for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text.startswith("단계 ")]
    assert len(steps) == 3
    for shape in slide.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.TEXT_BOX or shape.text_frame.text == "↓":
            continue
        box = _box(shape)
        card = next((card for card in map(_box, _cards(slide)) if card[0] <= box[0] and box[2] <= card[2] and card[1] <= box[1] < card[3]), None)
        if card is None:
            continue
        frame = shape.text_frame
        font_pt = frame.paragraphs[0].runs[0].font.size.pt
        per_line = (box[2] - box[0] - 0.2) * 72 / font_pt
        lines = sum(max(1, math.ceil(_em_width(p.text) / per_line)) for p in frame.paragraphs)
        assert frame.word_wrap is True and box[1] + 0.05 + lines * font_pt * 1.2 / 72 <= card[3] + 0.01, (frame.text, card)
    assert max(_box(card)[3] for card in _cards(slide) if _box(card)[0] >= 5.0 and _box(card)[1] < 4.4) <= 4.45


def test_fit_card_lines_marks_dropped_lines_with_an_ellipsis():
    from app.services.pptx.primitives import _fit_card_lines

    assert _fit_card_lines(["짧은 줄", "두번째 줄"], width=3.04, max_lines=1, font_size_pt=11) == ["짧은 줄…"]
    assert _fit_card_lines(["짧은 줄", "두번째 줄"], width=3.04, max_lines=2, font_size_pt=11) == ["짧은 줄", "두번째 줄"]
    long_line = "가" * 30
    fitted = _fit_card_lines([long_line, "다음 줄"], width=3.04, max_lines=2, font_size_pt=11)
    assert fitted[-1].endswith("…")
    assert "다음 줄" not in fitted
    assert _fit_card_lines(["**", "  "], width=3.04, max_lines=2, font_size_pt=11) == []


def test_cards_with_markup_only_titles_or_captions_do_not_fail():
    import base64

    from pptx import Presentation as NewPresentation

    from app.services.pptx.primitives import _add_card
    from app.services.pptx.visual_panels import _render_visual_image_asset

    deck = NewPresentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _add_card(slide, left=0.5, top=0.5, width=3.0, height=1.0, title="**", body=["본문"])
    png = base64.b64encode(
        bytes.fromhex(
            "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
            "1f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082"
        )
    ).decode()
    asset = {"media_type": "image/png", "content_base64": png, "visual_brief": ""}
    assert _render_visual_image_asset(
        slide, {"visual_brief": "**"}, asset, left=5.0, top=1.25, width=4.0, height=3.0
    )
