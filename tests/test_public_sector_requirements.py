"""Requirement ID scan and trace on synthetic RFPs and drafts in every supported format."""
from __future__ import annotations

from pathlib import Path

from docx import Document

from app.public_sector.models import RfpProfile
from app.public_sector.readers import read_document
from app.public_sector.requirements import (
    RequirementKey,
    find_ids,
    parse_requirement_id,
    scan_requirements,
    trace_requirements,
    trace_table_markdown,
)
from app.services.hwp_service import build_hwp

RFP_MD = """# [합성] 문서관리시스템 고도화 제안요청서

## 1. 사업 개요
수행사는 ISO-9001 인증을 보유해야 한다.

## 2. 요구사항
| 요구사항 고유번호 | 요구사항 명칭 |
|---|---|
| SFR-001 | 통합 검색 기능 |
| SFR-002 | 권한별 열람 제어 |
| PER-001 | 검색 응답 시간 |
| SER-001 | 개인정보 암호화 |
"""

DRAFT_MD = """# 기술 제안서

## Ⅱ. 기술 부문
### 1. 검색 기능
통합 검색은 SFR-001을 충족한다. 응답 시간은 PER-001 기준을 따른다.
### 2. 권한 관리
SFR-2 권한별 열람 제어를 제공한다. 추가로 SFR-009 감사 로그를 제안한다.
"""


def _profile(*ids: str) -> RfpProfile:
    return RfpProfile.model_validate({
        "schema_version": "publicdoc.rfp_profile.v1",
        "project": {"name": "[합성] 문서관리시스템 고도화"},
        "requirements": [{"id": value, "name": f"{value} 명칭"} for value in ids],
    })


def _write(tmp_path: Path, name: str, text: str) -> Path:
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


def test_ids_compare_by_prefix_and_number_in_ascii_and_korean_forms() -> None:
    found = find_ids("SFR-001과 SFR - 1, 기능 요구사항-003, COVID19는 아니다")

    assert [match.key for match in found] == [
        RequirementKey("SFR", 1), RequirementKey("SFR", 1), RequirementKey("기능요구사항", 3),
    ]
    assert parse_requirement_id("PER-010") == RequirementKey("PER", 10)
    assert parse_requirement_id("성능 기준") is None


def test_scan_lists_repeated_or_conventional_prefixes_with_names(tmp_path: Path) -> None:
    report = scan_requirements(read_document(_write(tmp_path, "rfp.md", RFP_MD)))
    candidates = {item["id"]: item for item in report["candidates"]}

    assert sorted(candidates) == ["PER-001", "SER-001", "SFR-001", "SFR-002"]
    assert candidates["SFR-002"]["name_candidate"] == "권한별 열람 제어"
    assert candidates["SER-001"]["category_guess"] == "보안"
    assert candidates["SFR-001"]["heading_path"].endswith("2. 요구사항")
    # ISO-9001 appears once with an unknown prefix, so it is not a requirement candidate.
    assert report["prefix_mentions"]["ISO"] == 1


def test_trace_reports_missing_unknown_and_writes_the_trace_table(tmp_path: Path) -> None:
    report = trace_requirements(
        _profile("SFR-001", "SFR-002", "PER-001", "SER-001"),
        read_document(_write(tmp_path, "draft.md", DRAFT_MD)),
    )

    assert report["summary"] == {"requirements": 4, "covered": 3, "missing": 1, "unknown_ids": 1}
    assert report["missing"] == ["SER-001"]
    assert [item["id"] for item in report["unknown_ids"]] == ["SFR-009"]
    assert report["passed"] is False
    sfr2 = next(row for row in report["rows"] if row["id"] == "SFR-002")
    assert sfr2["locations"][0]["heading_path"] == "기술 제안서 > Ⅱ. 기술 부문 > 2. 권한 관리"
    table = trace_table_markdown(report)
    assert "| SER-001 | 확인 필요 | SER-001 명칭 | - | 누락 |" in table


def test_trace_flags_duplicate_profile_ids(tmp_path: Path) -> None:
    report = trace_requirements(_profile("SFR-001", "SFR-1"), read_document(_write(tmp_path, "d.md", DRAFT_MD)))

    assert [issue["code"] for issue in report["profile_issues"]] == ["duplicate_id"]
    assert report["passed"] is False


def test_trace_reads_docx_headings_and_tables(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("기술 제안서", level=1)
    document.add_heading("검색 기능", level=2)
    document.add_paragraph("통합 검색으로 요구사항을 충족한다.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "SFR-001"
    table.rows[0].cells[1].text = "충족"
    path = tmp_path / "draft.docx"
    document.save(path)

    report = trace_requirements(_profile("SFR-001", "SER-001"), read_document(path))

    covered = next(row for row in report["rows"] if row["id"] == "SFR-001")
    assert covered["locations"][0]["heading_path"] == "기술 제안서 > 검색 기능"
    assert covered["locations"][0]["excerpt"] == "SFR-001 | 충족"
    assert report["missing"] == ["SER-001"]


def test_trace_reads_hwpx_written_by_the_exporter(tmp_path: Path) -> None:
    markdown = "# 기술 제안서\n\n## 보안 대책\n개인정보는 SER-001에 따라 암호화한다.\n"
    path = tmp_path / "draft.hwpx"
    path.write_bytes(build_hwp([{"doc_type": "tech_proposal", "markdown": markdown}], "[합성] 제안서"))

    report = trace_requirements(_profile("SER-001"), read_document(path))

    assert report["passed"] is True
    # The exporter's cover summary repeats body text, so the body heading is one of the locations.
    assert any(location["heading_path"].endswith("보안 대책") for location in report["rows"][0]["locations"])


def test_scan_strips_cell_labels_from_flattened_table_rows(tmp_path: Path) -> None:
    rows = (
        "# RFP\n\n## 요구사항 총괄표\n"
        "• 구분: 기능 / 요구사항 고유번호: SFR-001 / 요구사항 명칭: 통합 검색 / 우선순위: 상\n"
        "• 구분: 보안 / 요구사항 고유번호: SER-001 / 요구사항 명칭: 개인정보 암호화\n"
    )
    candidates = {c["id"]: c["name_candidate"] for c in scan_requirements(read_document(_write(tmp_path, "rfp.md", rows)))["candidates"]}

    assert candidates == {"SFR-001": "통합 검색", "SER-001": "개인정보 암호화"}


def test_bundle_parts_are_traced_as_one_document_in_order(tmp_path: Path) -> None:
    from app.public_sector.readers import read_documents

    first = _write(tmp_path, "part1.md", "# 사업 이해\n\n## 배경\nSFR-001 검색을 다룬다.\n")
    second = _write(tmp_path, "part2.md", "# 기술 제안\n\n## 보안 대책\nSER-001 암호화를 적용한다.\n")

    document = read_documents([first, second])
    report = trace_requirements(_profile("SFR-001", "SER-001"), document)

    assert report["passed"] is True
    assert document.path == f"{first}, {second}"
    assert report["rows"][1]["locations"][0]["heading_path"] == "기술 제안 > 보안 대책"


def test_trace_table_drops_the_title_every_location_shares(tmp_path: Path) -> None:
    single = trace_requirements(_profile("SFR-001", "SER-001"), read_document(_write(tmp_path, "d.md", "# 제안서\n\n## 검색\nSFR-001\n\n## 보안\nSER-001\n")))
    table = trace_table_markdown(single)

    assert "| SFR-001 | 확인 필요 | SFR-001 명칭 | 검색 | 대응 |" in table
    assert "제안서 >" not in table

    from app.public_sector.readers import read_documents

    parts = read_documents([_write(tmp_path, "a.md", "# 사업 이해\n\n## 배경\nSFR-001\n"), _write(tmp_path, "b.md", "# 기술 제안\n\n## 보안\nSER-001\n")])
    assert "| 기술 제안 > 보안 | 대응 |" in trace_table_markdown(trace_requirements(_profile("SFR-001", "SER-001"), parts))
