"""Portable applicability evidence, with no provider or user-data access."""
import copy
import hashlib
import html
import io
import json
import zipfile
from dataclasses import replace
from unittest.mock import Mock
from uuid import uuid4

import pytest

from app.services.procurement_decision_package.package_builder import (
    _package_checklist_status,
    build_decision_package_from_record,
)
from app.services.procurement_decision_package.review_packet import (
    build_project_procurement_review_packet,
    verify_bound_procurement_packet,
    verify_procurement_review_packet,
)
from app.services.procurement_decision_package.review_receipt import (
    build_pending_procurement_review_receipt, record_procurement_review_decision,
    validate_procurement_review_receipt,
)
from app.services.procurement_decision_package.reviewed_package import (
    build_procurement_reviewed_package, verify_procurement_reviewed_package,
    REVIEWED_PACKAGE_PACKET_NAME, REVIEWED_PACKAGE_RECEIPT_NAME,
)
from app.storage.state_backend import LocalStateBackend
from app.services.procurement_review_evidence import (
    validate_persisted_procurement_review_packet,
)
from app.storage.procurement_review_store import (
    ProcurementReviewStore, ProcurementReviewStoreError,
)
from tests.test_procurement_opportunity_bindings import _capture, _entry
from tests.test_procurement_review_store import _MemoryS3Client, _s3_backend


@pytest.fixture
def context(tmp_path):
    backend = LocalStateBackend(tmp_path)
    entry = _entry(backend)
    binding = _capture(entry, backend)
    source_hash = hashlib.sha256(json.dumps(
        [item.model_dump(mode="json") for item in binding.snapshots],
        ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False,
    ).encode()).hexdigest()
    quote = '<script>alert("x")</script> | **bold** 😀\n```\n[link](javascript:x)'
    source = dict(snapshot_id=binding.snapshots[0].snapshot_id,
                  snapshot_sha256=binding.snapshots[0].sha256,
                  source_set_sha256=source_hash, quote_start=3,
                  quote_end=3 + len(quote), quote=quote)
    annotation = dict(annotation_id=str(uuid4()), applicability="not_applicable",
                      rationale="Human review only", **source,
                      created_at="2026-09-22T01:00:00Z",
                      expected_decision_revision=1, operation_id=str(uuid4()))
    row = dict(requirement_id=str(uuid4()), title="Local requirement",
               category="domain_capability_fit", **source,
               created_at="2026-09-22T00:00:00Z", annotations=[annotation],
               applicability="not_applicable", stale=False)
    projection = dict(schema_version="procurement.requirement_applicability.v1",
                      decision_id=binding.decision_id, decision_revision=2,
                      source_set_sha256=source_hash, requirements=[row])
    return entry.record, binding, projection


def _build(context):
    record, binding, projection = context
    return build_project_procurement_review_packet(
        record, reviewer_owner="reviewer", source_binding=binding,
        requirement_applicability=projection,
    )


def _entries(content):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


@pytest.fixture(params=["local", "fake-s3"])
def persisted_packet(context, tmp_path, request):
    backend = (
        LocalStateBackend(tmp_path / "reviews")
        if request.param == "local" else _s3_backend(_MemoryS3Client())
    )
    packet = _build(context)
    receipt = build_pending_procurement_review_receipt(packet.content)
    store = ProcurementReviewStore(backend=backend)
    record, _ = store.prepare(
        tenant_id="alpha", project_id="project-a", packet_content=packet.content,
        receipt=receipt, prepared_at="2026-09-28T00:00:00Z",
    )
    return backend, record, packet.content


@pytest.mark.parametrize("operation", ["evidence", "filter", "filter-with-validator"])
def test_packet_read_avoids_duplicate_docx_rebuilds(
    persisted_packet, context, monkeypatch, operation,
):
    from app.services.procurement_decision_package import applicability

    backend, record, content = persisted_packet
    rebuild = Mock(wraps=applicability.build_applicability_docx)
    monkeypatch.setattr(applicability, "build_applicability_docx", rebuild)
    callback = Mock(wraps=validate_persisted_procurement_review_packet)
    with_callback = operation == "filter-with-validator"
    store = ProcurementReviewStore(
        backend=backend, packet_evidence_validator=callback if with_callback else None,
    )
    per_read = {"evidence": 1, "filter": 2, "filter-with-validator": 3}[operation]
    for read_number in (1, 2):
        if operation == "evidence":
            validate_persisted_procurement_review_packet(record, content)
        else:
            assert store.filter_by_decision(
                [record], tenant_id="alpha", project_id="project-a",
                decision_id=context[1].decision_id,
            ) == [record]
        # A later read still verifies bytes; no process-wide content cache.
        assert rebuild.call_count == read_number * per_read
        assert callback.call_count == (read_number if with_callback else 0)


@pytest.mark.parametrize("field,value", [
    ("tenant_id", "other"), ("project_id", "other"),
    ("reviewer", "other"), ("package_id", "other"),
])
def test_persisted_packet_rejects_scope_and_record_drift(persisted_packet, field, value):
    _, record, content = persisted_packet
    with pytest.raises(ValueError):
        validate_persisted_procurement_review_packet(replace(record, **{field: value}), content)


@pytest.mark.parametrize("failure", ["changed-packet", "validator-rejected"])
def test_decision_filter_rechecks_evidence_after_success(persisted_packet, context, failure):
    backend, record, _ = persisted_packet
    callback = Mock(wraps=validate_persisted_procurement_review_packet)
    store = ProcurementReviewStore(backend=backend, packet_evidence_validator=callback)
    scope = dict(tenant_id="alpha", project_id="project-a", decision_id=context[1].decision_id)
    assert store.filter_by_decision([record], **scope) == [record]
    prefix = f"tenants/alpha/procurement_reviews/project-a/{record.packet_sha256}"
    if failure == "changed-packet":
        backend.write_bytes(f"{prefix}/packet.zip", b"tampered after successful read")
    else:
        callback.side_effect = ValueError("additional evidence rejected")
    before = {path: backend.read_bytes(path) for path in backend.list_prefix(prefix)}
    with pytest.raises(ProcurementReviewStoreError):
        store.filter_by_decision([record], **scope)
    assert {path: backend.read_bytes(path) for path in backend.list_prefix(prefix)} == before
    assert callback.call_count == (1 if failure == "changed-packet" else 2)


def test_v3_roundtrip_and_unchanged_parent_evaluation(context):
    record, binding, projection = context
    before = record.model_dump(mode="json")
    legacy = build_decision_package_from_record(record, reviewer_owner="reviewer")
    packet = _build(context)
    assert packet.manifest["schema_version"] == "decisiondoc.procurement_review_packet.v3"
    assert packet.verification["source_bytes_verified"] is False
    assert packet.verification["operational_approval"] is False
    entries = _entries(packet.content)
    doc = json.loads(entries["decision_package.json"])
    assert doc["schema_purpose"] == "procurement_decision_package.v2"
    assert doc["package"]["requirement_applicability"] == projection
    for field in ("bid_readiness_checklist", "soft_fit_score", "recommendation", "hard_filters"):
        assert doc["package"][field] == legacy["package"][field]
    assert "requirement_applicability.docx" in entries
    assert record.model_dump(mode="json") == before
    assert _build(context).content == packet.content


def test_projection_requires_binding(context):
    record, _, projection = context
    for builder in (build_decision_package_from_record, build_project_procurement_review_packet):
        with pytest.raises(ValueError, match="binding"):
            builder(record, reviewer_owner="reviewer", requirement_applicability=projection)


@pytest.mark.parametrize("status", ["not_applicable", "needs_review", "future", "", None])
def test_unknown_parent_status_rejected(status):
    with pytest.raises(ValueError):
        _package_checklist_status(status)


def _json(value):
    return (json.dumps(value, ensure_ascii=False, indent=2) + "\n").encode()


def _rewrite(entries):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return output.getvalue()


def _rehash(entries):
    manifest = json.loads(entries["packet_manifest.json"])
    manifest["artifacts"] = [dict(path=name, size_bytes=len(content),
                                   sha256=hashlib.sha256(content).hexdigest())
                             for name, content in entries.items() if name != "packet_manifest.json"]
    manifest["artifact_count"] = len(manifest["artifacts"])
    entries["packet_manifest.json"] = _json(manifest)
    return _rewrite(entries)


@pytest.mark.parametrize("path,value", [
    (("actor_id",), "secret"),
    (("decision_id",), str(uuid4())),
    (("decision_revision",), True),
    (("decision_revision",), 3),
    (("source_set_sha256",), "f" * 64),
    (("schema_version",), "procurement.requirement_applicability.v2"),
    (("requirements", 0, "created_by_actor_id"), "secret"),
    (("requirements", 0, "requirement_id"), "not-uuid"),
    (("requirements", 0, "applicability"), "ready"),
    (("requirements", 0, "applicability"), "unknown"),
    (("requirements", 0, "stale"), True),
    (("requirements", 0, "stale"), 0),
    (("requirements", 0, "quote_start"), True),
    (("requirements", 0, "quote_end"), 1000),
    (("requirements", 0, "category"), "invented"),
    (("requirements", 0, "annotations", 0, "actor_id"), "secret"),
    (("requirements", 0, "annotations", 0, "snapshot_sha256"), "b" * 64),
    (("requirements", 0, "annotations", 0, "source_set_sha256"), "b" * 64),
    (("requirements", 0, "annotations", 0, "expected_decision_revision"), 2),
    (("requirements", 0, "annotations", 0, "operation_id"), "invalid"),
    (("requirements", 0, "annotations", 0, "rationale"), "   "),
    (("requirements", 0, "annotations", 0, "applicability"), "future"),
])
def test_strict_projection_rejects_malicious_fields(context, path, value):
    projection = context[2]
    target = projection
    for key in path[:-1]:
        target = target[key]
    target[path[-1]] = value
    with pytest.raises(ValueError):
        _build(context)


def test_empty_projection_and_no_history_are_unknown(context):
    row = context[2]["requirements"][0]
    row["annotations"] = []
    row["applicability"] = "unknown"
    _build(context)
    context[2]["requirements"] = []
    packet = _build(context)
    assert packet.verification["schema_version"].endswith(".v3")


def test_stale_retains_history_and_counts_separate(context):
    from app.services.procurement_decision_package.applicability import applicability_counts
    row = context[2]["requirements"][0]
    row["source_set_sha256"] = "a" * 64
    row["annotations"][0]["source_set_sha256"] = "a" * 64
    row.update(stale=True, applicability="unknown")
    packet = _build(context)
    counts = applicability_counts(context[2])
    assert counts == dict(total=1, applies=0, not_applicable=0, unknown=1, stale=1)
    entries = _entries(packet.content)
    assert '근거 변경(stale) 1' in entries["bid_readiness_checklist.md"].decode()
    assert 'not_applicable' in html.unescape(entries["bid_readiness_checklist.md"].decode())
    assert json.loads(entries["decision_package.json"])["package"]["requirement_applicability"] == context[2]


def test_history_latest_retraction_and_identity_checks(context):
    record, binding, projection = context
    binding = binding.model_copy(update={"decision_revision": 4})
    projection["decision_revision"] = 4
    row = projection["requirements"][0]
    original = copy.deepcopy(row["annotations"][0])
    retraction = {**original, "annotation_id": str(uuid4()), "operation_id": str(uuid4()),
                  "expected_decision_revision": 3, "applicability": "unknown", "rationale": "Withdrawn"}
    row["annotations"].append(retraction)
    row["applicability"] = "unknown"
    _build((record, binding, projection))
    for field in ("annotation_id", "operation_id"):
        prior = retraction[field]
        retraction[field] = original[field]
        with pytest.raises(ValueError):
            _build((record, binding, projection))
        retraction[field] = prior
    row["annotations"].reverse()
    with pytest.raises(ValueError):
        _build((record, binding, projection))


@pytest.mark.parametrize("artifact", ["bid_readiness_checklist.md", "procurement_review.html",
                                      "requirement_applicability.docx"])
def test_rehashed_rendered_content_tamper_is_rejected(context, artifact):
    entries = _entries(_build(context).content)
    if artifact.endswith("docx"):
        inner = _entries(entries[artifact])
        inner["word/document.xml"] = inner["word/document.xml"].replace(b"not_applicable", b"applies")
        entries[artifact] = _rewrite(inner)
    else:
        entries[artifact] = entries[artifact].replace(b"not_applicable", b"applies")
    with pytest.raises(ValueError, match="content mismatch"):
        verify_procurement_review_packet(_rehash(entries))


@pytest.mark.parametrize("mutation", ["missing-docx", "missing-section", "downgrade-v2", "downgrade-v1", "unknown"])
def test_packet_versions_cannot_drop_applicability(context, mutation):
    entries = _entries(_build(context).content)
    doc = json.loads(entries["decision_package.json"])
    manifest = json.loads(entries["packet_manifest.json"])
    if mutation == "missing-docx":
        del entries["requirement_applicability.docx"]
    elif mutation == "missing-section":
        del doc["package"]["requirement_applicability"]
    elif mutation.startswith("downgrade"):
        manifest["schema_version"] = "decisiondoc.procurement_review_packet." + mutation[-2:]
        del entries["requirement_applicability.docx"]
        if mutation.endswith("v1"):
            del manifest["source_binding"]
    else:
        manifest["schema_version"] = "decisiondoc.procurement_review_packet.v99"
    entries["decision_package.json"] = _json(doc)
    entries["packet_manifest.json"] = _json(manifest)
    content = _rehash(entries)
    with pytest.raises(ValueError):
        verify_procurement_review_packet(content)
    with pytest.raises(ValueError):
        verify_bound_procurement_packet(content, expected_tenant_id=context[1].tenant_id,
                                         expected_project_id=context[1].project_id)


def test_completed_review_binds_full_v3_and_preserves_source_bytes(context):
    packet = _build(context)
    pending = build_pending_procurement_review_receipt(packet.content)
    receipt = record_procurement_review_decision(
        pending, packet.content, reviewer="reviewer", decision="accepted",
        rationale="Review only", reviewed_at="2026-09-22T12:00:00Z",
    )
    receipt_bytes = _json(receipt)
    completed, manifest = build_procurement_reviewed_package(packet.content, receipt,
                                                             receipt_content=receipt_bytes)
    assert manifest["source"]["packet_schema_version"].endswith(".v3")
    result = verify_procurement_reviewed_package(completed)
    assert result["source_bytes_verified"] is False
    assert result["operational_approval"] is False
    entries = _entries(completed)
    assert entries[REVIEWED_PACKAGE_PACKET_NAME] == packet.content
    assert entries[REVIEWED_PACKAGE_RECEIPT_NAME] == receipt_bytes
    context[2]["requirements"][0]["annotations"][0]["rationale"] = "Changed reason"
    changed = _build(context)
    with pytest.raises(ValueError, match="packet_sha256"):
        validate_procurement_review_receipt(receipt, changed.content)
    assert verify_procurement_reviewed_package(completed)["package_verified"] is True


def test_legacy_versions_and_inventory_unchanged_after_v3(context):
    from app.services.procurement_decision_package.package_constants import INCLUDED_ARTIFACT_ORDER
    record, binding, _ = context
    first = build_project_procurement_review_packet(record, reviewer_owner="reviewer")
    second = build_project_procurement_review_packet(record, reviewer_owner="reviewer", source_binding=binding)
    _build(context)
    assert build_project_procurement_review_packet(record, reviewer_owner="reviewer").content == first.content
    assert build_project_procurement_review_packet(record, reviewer_owner="reviewer", source_binding=binding).content == second.content
    a, b = _entries(first.content), _entries(second.content)
    assert list(a) == list(b) == [*INCLUDED_ARTIFACT_ORDER, "packet_manifest.json"]
    assert all(a[name] == b[name] for name in INCLUDED_ARTIFACT_ORDER)
    assert "requirement_applicability.docx" not in INCLUDED_ARTIFACT_ORDER


def test_literal_quotes_escape_html_markdown_and_survive_docx(context):
    from docx import Document
    from app.services.procurement_decision_package.applicability import applicability_lines
    entries = _entries(_build(context).content)
    assert b"<script>" not in entries["procurement_review.html"]
    assert b"&lt;script&gt;" in entries["procurement_review.html"]
    assert b"created_by_actor_id" not in entries["decision_package.json"]
    assert b'"actor_id"' not in entries["decision_package.json"]
    lines = applicability_lines(context[2])
    markdown = entries["bid_readiness_checklist.md"].decode()
    quote = context[2]["requirements"][0]["quote"]
    assert "\n".join("    " + line for line in quote.split("\n")) in markdown
    doc = Document(io.BytesIO(entries["requirement_applicability.docx"]))
    assert [p.text for p in doc.paragraphs][-len(lines):] == lines
    assert any(p.style.name == "Heading 2" and p.text.startswith("요구사항 1.") for p in doc.paragraphs)
    assert any(p.style.name == "Heading 3" and p.text.startswith("이력 1.") for p in doc.paragraphs)


@pytest.mark.parametrize("field,value", [("snapshot_id", "absent"), ("snapshot_sha256", "c" * 64)])
def test_current_requirement_must_resolve_snapshot_in_binding(context, field, value):
    row = context[2]["requirements"][0]
    row[field] = value
    row["annotations"][0][field] = value
    with pytest.raises(ValueError, match="snapshot is not in source_binding"):
        _build(context)


def test_rehashed_section_semantic_tamper_rejected(context):
    entries = _entries(_build(context).content)
    doc = json.loads(entries["decision_package.json"])
    row = doc["package"]["requirement_applicability"]["requirements"][0]
    row["applicability"] = "applies"
    entries["decision_package.json"] = _json(doc)
    with pytest.raises(ValueError, match="effective applicability"):
        verify_procurement_review_packet(_rehash(entries))


def test_crlf_source_quotes_remain_exact_in_portable_projection(context):
    from docx import Document
    from docx.oxml.ns import qn

    quote = "공고 A\r\n😀 원문\r\n마지막 줄"
    row = context[2]["requirements"][0]
    for item in (row, row["annotations"][0]):
        item.update(quote=quote, quote_start=0, quote_end=len(quote))
    packet = _build(context)
    entries = _entries(packet.content)
    assert json.loads(entries["decision_package.json"])["package"]["requirement_applicability"]["requirements"][0]["quote"] == quote
    doc = Document(io.BytesIO(entries["requirement_applicability.docx"]))
    assert doc.paragraphs[0].style.name == "Title"
    assert str(doc.paragraphs[0].runs[0].font.color.rgb) == "000000"
    assert not doc.styles["Title"].element.findall(".//" + qn("w:pBdr"))
    assert quote.replace("\r\n", "\n") in [paragraph.text for paragraph in doc.paragraphs]
