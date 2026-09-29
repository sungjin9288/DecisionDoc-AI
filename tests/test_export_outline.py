from app.services.export_outline import presentation_points, summarize_export_docs
from app.services.export_labels import humanize_doc_type


def test_core_document_export_labels_preserve_ids_and_existing_fallbacks():
    labels = {
        "adr": "기술 의사결정 기록 (ADR)",
        "onepager": "한 페이지 요약",
        "eval_plan": "평가 계획",
        "ops_checklist": "운영 체크리스트",
    }
    docs = [{"doc_type": key, "markdown": "# Title\n\nBody"} for key in labels]
    summaries = summarize_export_docs(docs)
    for doc, summary in zip(docs, summaries):
        assert summary["label"] == labels[doc["doc_type"]]
        assert labels[doc["doc_type"]] == humanize_doc_type(doc["doc_type"])
    assert [doc["doc_type"] for doc in docs] == list(labels)
    assert humanize_doc_type("business_understanding") == "사업 이해"
    assert humanize_doc_type("future_document") == "Future Document"
    assert humanize_doc_type("") == "문서"


def test_presentation_points_split_long_sentence_into_clauses() -> None:
    text = (
        "본 제안은 핵심 정책 목표를 공공기관이 실제 운영 KPI로 관리할 수 있도록 데이터 통합, "
        "AI 분석, 운영 대시보드를 하나의 사업 범위로 묶은 안입니다."
    )
    points = presentation_points(text, max_len=48, max_points=4)
    assert len(points) >= 2
    assert any("AI 분석 · 운영 대시보드" in point for point in points)
    assert all(len(point) <= 48 for point in points)
    assert "AI 분석" not in points


def test_presentation_points_merges_short_enumeration_fragments() -> None:
    text = (
        "수행계획서는 계약 범위, 일정, 산출물, 투입 인력, 승인 게이트를 하나의 실행 문서로 정리한 결과물입니다."
    )
    points = presentation_points(text, max_len=40, max_points=4)
    assert any("계약 범위 · 일정 · 산출물" in point for point in points)
    assert not any(point == "일정" for point in points)


def test_summarize_export_docs_exposes_short_ppt_lead() -> None:
    docs = [
        {
            "doc_type": "business_understanding",
            "markdown": (
                "# 사업 이해\n\n"
                "첫 문장은 발표자료용 요약으로 충분히 짧아야 합니다. "
                "두 번째 문장은 문서형 상세 설명입니다."
            ),
        }
    ]
    summary = summarize_export_docs(docs)[0]
    assert summary["ppt_lead"] == "첫 문장은 발표자료용 요약으로 충분히 짧아야 합니다."
    assert "두 번째 문장" not in summary["ppt_lead"]


def test_summarize_export_docs_exposes_structured_section_and_metric_items() -> None:
    docs = [
        {
            "doc_type": "business_understanding",
            "markdown": (
                "# 사업 이해\n\n"
                "제안의 핵심 요약입니다.\n\n"
                "## 제안 요약\n\n"
                "| 항목 | 내용 |\n| --- | --- |\n| KPI | 운영 정렬 |\n\n"
                "## 사업 배경\n\n"
                "- 정책 배경\n"
            ),
        }
    ]
    summary = summarize_export_docs(docs)[0]
    assert summary["section_items"] == ["제안 요약", "사업 배경"]
    assert summary["metric_items"] == ["표 1개", "목록 1개"]


def test_presentation_points_keep_ordinal_markers_with_their_text() -> None:
    assert presentation_points("1. 번호 목록 하나", max_len=78, max_points=6) == ["1. 번호 목록 하나"]
    assert presentation_points("2.1. 세부 단계입니다. 다음 문장입니다.", max_len=78, max_points=6) == [
        "2.1. 세부 단계입니다.",
        "다음 문장입니다.",
    ]


def test_pptx_numbered_paragraphs_stay_on_one_bullet() -> None:
    from io import BytesIO

    from pptx import Presentation

    from app.services.pptx_service import build_pptx_from_docs

    markdown = "# 실행\n\n## 실행 항목\n\n- 첫째\n- 둘째\n- 셋째\n1. 번호 목록 하나\n2. 번호 목록 둘\n"
    deck = Presentation(BytesIO(build_pptx_from_docs([{"doc_type": "adr", "markdown": markdown}], "번호")))
    paragraphs = [
        paragraph.text
        for slide in deck.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
    ]

    assert "1. 번호 목록 하나" in paragraphs
    assert "2. 번호 목록 둘" in paragraphs
    assert "1." not in paragraphs and "2." not in paragraphs
